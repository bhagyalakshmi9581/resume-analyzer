from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from flask import Flask, render_template, request
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from analyzer import analyze_resume

app = Flask(__name__)


def extract_pdf_text(pdf_data: bytes) -> str:
    """Extract text from all pages in a PDF document."""
    reader = PdfReader(BytesIO(pdf_data))
    text = "\n".join(
        page_text
        for page in reader.pages
        if (page_text := page.extract_text())
    ).strip()
    if not text:
        raise ValueError("The uploaded PDF does not contain extractable text.")
    return text


def extract_docx_text(docx_data: bytes) -> str:
    """Extract text from paragraphs and tables in a Word document."""
    document = Document(BytesIO(docx_data))
    text_parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    text_parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    )
    text = "\n".join(text_parts).strip()
    if not text:
        raise ValueError("The uploaded Word document contains no extractable text.")
    return text


@app.route("/", methods=["GET", "POST"])
def home() -> str:
    result = None
    error = None

    if request.method == "POST":
        uploaded_file = request.files.get("resume_file")
        if uploaded_file and uploaded_file.filename:
            file_extension = uploaded_file.filename.lower().rsplit(".", 1)[-1]
            if file_extension == "pdf":
                extract_text = extract_pdf_text
                file_type = "PDF"
                extraction_errors = (PdfReadError, ValueError)
            elif file_extension == "docx":
                extract_text = extract_docx_text
                file_type = "Word"
                extraction_errors = (BadZipFile, PackageNotFoundError, ValueError)
            else:
                error = "Please upload a PDF or Word (.docx) file."

            if not error:
                try:
                    resume = extract_text(uploaded_file.read())
                except extraction_errors:
                    error = (
                        f"Unable to extract text from the uploaded {file_type} file. "
                        "Please use a readable document with extractable text."
                    )
        else:
            resume = request.form.get("resume", "")

        if not error:
            job = request.form.get("job", "")
            result = analyze_resume(resume, job)

    return render_template(
        "index.html",
        result=result,
        error=error,
    )

if __name__ == "__main__":
    app.run(debug=True)
