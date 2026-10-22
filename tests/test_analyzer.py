from io import BytesIO
from zipfile import BadZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.errors import PdfReadError

import app as resume_app
from analyzer import analyze_resume
from app import app, extract_docx_text, extract_pdf_text


def test_analyze_resume_reports_missing_skills():
    result = analyze_resume("Python Flask", "Python Flask SQL")

    assert result["score"] == 67
    assert result["matched_skills"] == ["flask", "python"]
    assert result["missing_skills"] == ["sql"]
    assert result["suggestions"] == ["Consider adding sql"]


def test_analyze_resume_handles_empty_job_description():
    result = analyze_resume("Python Flask", "")

    assert result == {
        "score": 0,
        "matched_skills": [],
        "missing_skills": [],
        "suggestions": [],
    }


def test_analyze_resume_excludes_common_words_from_missing_skills():
    result = analyze_resume(
        "Python Flask are be can",
        "It is taking Python Flask and are be can also SQL",
    )

    assert result["score"] == 67
    assert result["matched_skills"] == ["flask", "python"]
    assert result["missing_skills"] == ["sql"]


def test_results_page_displays_missing_skill_count_and_list():
    response = app.test_client().post(
        "/",
        data={"resume": "Python", "job": "Python Flask"},
    )

    assert response.status_code == 200
    assert b"Missing Skills" in response.data
    assert b"(1)" in response.data
    assert b"<li>flask</li>" in response.data


def test_results_page_displays_message_when_no_skills_are_missing():
    response = app.test_client().post(
        "/",
        data={"resume": "Python Flask", "job": "Python Flask"},
    )

    assert b"No missing skills identified." in response.data


def test_extract_pdf_text_combines_page_text(monkeypatch):
    class Page:
        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

    class Reader:
        def __init__(self, _pdf_file):
            self.pages = [Page("Python"), Page(None), Page("Flask")]

    monkeypatch.setattr(resume_app, "PdfReader", Reader)

    assert extract_pdf_text(b"PDF data") == "Python\nFlask"


def test_extract_pdf_text_rejects_pdf_without_text():
    pdf_file = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(pdf_file)

    with pytest.raises(ValueError, match="does not contain extractable text"):
        extract_pdf_text(pdf_file.getvalue())


def test_extract_docx_text_extracts_paragraphs_and_table_cells():
    document_file = BytesIO()
    document = Document()
    document.add_paragraph("Python developer")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Flask experience"
    document.save(document_file)

    assert extract_docx_text(document_file.getvalue()) == (
        "Python developer\nFlask experience"
    )


def test_extract_docx_text_rejects_document_without_text():
    document_file = BytesIO()
    Document().save(document_file)

    with pytest.raises(ValueError, match="contains no extractable text"):
        extract_docx_text(document_file.getvalue())


def test_uploading_pdf_resume_analyzes_extracted_text(monkeypatch):
    monkeypatch.setattr(resume_app, "extract_pdf_text", lambda _data: "Python Flask")

    response = app.test_client().post(
        "/",
        data={
            "resume_file": (BytesIO(b"PDF data"), "resume.pdf"),
            "job": "Python Flask SQL",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert b"Match Score: 67%" in response.data
    assert b"<li>sql</li>" in response.data


def test_uploading_docx_resume_analyzes_extracted_text(monkeypatch):
    monkeypatch.setattr(
        resume_app,
        "extract_docx_text",
        lambda _data: "Python Flask",
    )

    response = app.test_client().post(
        "/",
        data={
            "resume_file": (BytesIO(b"Word data"), "resume.docx"),
            "job": "Python Flask SQL",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert b"Match Score: 67%" in response.data
    assert b"<li>sql</li>" in response.data


def test_uploading_unsupported_file_shows_validation_error():
    response = app.test_client().post(
        "/",
        data={
            "resume_file": (BytesIO(b"not a PDF"), "resume.txt"),
            "job": "Python",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert b"Please upload a PDF or Word (.docx) file." in response.data


def test_uploading_unreadable_pdf_shows_error(monkeypatch):
    def raise_pdf_error(_data):
        raise PdfReadError("Invalid PDF")

    monkeypatch.setattr(resume_app, "extract_pdf_text", raise_pdf_error)

    response = app.test_client().post(
        "/",
        data={
            "resume_file": (BytesIO(b"not a PDF"), "resume.pdf"),
            "job": "Python",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert b"Unable to extract text from the uploaded PDF file." in response.data


def test_uploading_unreadable_docx_shows_error(monkeypatch):
    def raise_bad_zip(_data):
        raise BadZipFile("Invalid Word document")

    monkeypatch.setattr(resume_app, "extract_docx_text", raise_bad_zip)

    response = app.test_client().post(
        "/",
        data={
            "resume_file": (BytesIO(b"not a DOCX"), "resume.docx"),
            "job": "Python",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert b"Unable to extract text from the uploaded Word file." in response.data